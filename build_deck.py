# -*- coding: utf-8 -*-
"""
Génère le diaporama de soutenance (Master 2) de Maxime LOUSTALET.
Sujet : L'intégration de l'IA dans les PME - le cas d'Arla Groupe.
Design system cohérent + notes du présentateur (script oral) sur chaque slide.
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import copy

# ---------------------------------------------------------------- PALETTE
NAVY     = RGBColor(0x0B, 0x2A, 0x4A)   # bleu nuit - primaire
NAVY2    = RGBColor(0x12, 0x3A, 0x61)   # bleu profond secondaire
STEEL    = RGBColor(0x3D, 0x6A, 0x8E)   # bleu acier
ACCENT   = RGBColor(0xE6, 0xA3, 0x39)   # ambre (accent / bâtiment)
ACCENT_D = RGBColor(0xAF, 0x71, 0x12)   # ambre foncé
LIGHT    = RGBColor(0xED, 0xF2, 0xF6)   # fond clair (mist)
CARD     = RGBColor(0xFF, 0xFF, 0xFF)
CARD_ALT = RGBColor(0xF7, 0xF2, 0xE8)   # crème chaud
PAPER    = RGBColor(0xFC, 0xFB, 0xF8)   # papier chaud (fond diapo)
TEXTD    = RGBColor(0x18, 0x24, 0x3C)   # texte foncé
TEXTG    = RGBColor(0x54, 0x60, 0x7A)   # texte gris
WHITE    = RGBColor(0xFF, 0xFF, 0xFF)
LINE_G   = RGBColor(0xE6, 0xDF, 0xD3)   # filet chaud

SERIF = "Georgia"
HEAD = "Georgia"   # titres / numéraux — allure éditoriale
BODY = "Calibri"   # texte courant / labels

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]

# ---------------------------------------------------------------- HELPERS
def slide():
    return prs.slides.add_slide(BLANK)

def _noline(sp):
    sp.line.fill.background()

def rect(s, x, y, w, h, color=None, line=None, lw=1.0, shape=MSO_SHAPE.RECTANGLE):
    sp = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if color is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = color
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line; sp.line.width = Pt(lw)
    sp.shadow.inherit = False
    return sp

def bg(s, color):
    rect(s, 0, 0, 13.333, 7.5, color)

def txt(s, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
        wrap=True, space_after=6, line_spacing=1.0):
    """runs: list of paragraphs; each paragraph = list of (text, size, bold, color, font, italic)."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        p.space_before = Pt(0)
        p.line_spacing = line_spacing
        for (t, size, bold, color, font, *rest) in para:
            italic = rest[0] if rest else False
            r = p.add_run(); r.text = t
            r.font.size = Pt(size); r.font.bold = bold
            r.font.color.rgb = color; r.font.name = font
            r.font.italic = italic
    return tb

def bullets(s, x, y, w, h, items, size=15, color=TEXTD, gap=9, lh=1.06,
            marker_color=ACCENT, bold_lead=True):
    """items: list of (lead, rest) OR plain string. Renders a subtle bullet."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap); p.space_before = Pt(0); p.line_spacing = lh
        rm = p.add_run(); rm.text = "▪  "
        rm.font.size = Pt(size); rm.font.color.rgb = marker_color; rm.font.name = BODY; rm.font.bold = True
        if isinstance(it, tuple):
            lead, rest = it
            r1 = p.add_run(); r1.text = lead
            r1.font.size = Pt(size); r1.font.bold = True; r1.font.color.rgb = color; r1.font.name = BODY
            if rest:
                r2 = p.add_run(); r2.text = rest
                r2.font.size = Pt(size); r2.font.bold = False; r2.font.color.rgb = TEXTG; r2.font.name = BODY
        else:
            r = p.add_run(); r.text = it
            r.font.size = Pt(size); r.font.color.rgb = color; r.font.name = BODY
    return tb

def header(s, kicker, title, num=None):
    """Top zone for content slides."""
    bg(s, PAPER)
    rect(s, 0, 0, 13.333, 1.5, PAPER)
    # left accent tab
    rect(s, 0, 0.0, 0.16, 1.5, ACCENT)
    txt(s, 0.75, 0.34, 11.5, 0.4,
        [[(kicker.upper(), 12, True, STEEL, BODY)]], space_after=0)
    txt(s, 0.75, 0.63, 11.9, 0.7,
        [[(title, 27, True, NAVY, HEAD)]], space_after=0)
    rect(s, 0.78, 1.32, 0.85, 0.055, ACCENT)
    if num:
        txt(s, 12.1, 0.30, 1.0, 0.6, [[(num, 40, True, RGBColor(0xE7,0xEC,0xF2), HEAD)]],
            align=PP_ALIGN.RIGHT, space_after=0)

def footer(s, page, dark=False):
    c = RGBColor(0x9A,0xA6,0xB8) if not dark else RGBColor(0x8A,0x9C,0xB5)
    txt(s, 0.75, 7.03, 8.0, 0.35,
        [[("Maxime LOUSTALET  ·  Soutenance de mémoire  ·  Eklore-ed School of Management  ·  2025–2026", 9, False, c, BODY)]],
        space_after=0)
    txt(s, 11.9, 7.03, 0.7, 0.35, [[(str(page), 9, True, c, BODY)]],
        align=PP_ALIGN.RIGHT, space_after=0)

def card(s, x, y, w, h, fill=CARD, line=LINE_G, lw=1.0):
    sp = rect(s, x, y, w, h, fill, line=line, lw=lw, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    try:
        sp.adjustments[0] = 0.045
    except Exception:
        pass
    return sp

def notes(s, text):
    s.notes_slide.notes_text_frame.text = text

# ================================================================ 1. TITRE
s = slide()
bg(s, NAVY)
# geometric accent band
rect(s, 0, 0, 13.333, 0.28, ACCENT)
rect(s, 0, 6.55, 13.333, 0.95, NAVY2)
rect(s, 0, 6.55, 13.333, 0.05, ACCENT)
# subtle side blocks
rect(s, 10.7, 1.2, 2.2, 0.05, STEEL)
txt(s, 0.9, 1.15, 11.5, 0.5,
    [[("SOUTENANCE DE MÉMOIRE DE FIN D'ÉTUDES  ·  MASTER 2", 13.5, True, ACCENT, HEAD)]], space_after=0)
txt(s, 0.9, 1.95, 11.6, 2.3,
    [[("L'intégration de l'intelligence", 41, True, WHITE, HEAD)],
     [("artificielle dans les PME", 41, True, WHITE, HEAD)]],
    space_after=2, line_spacing=1.0)
txt(s, 0.9, 3.95, 11.4, 0.9,
    [[("Entre opportunités perçues et freins organisationnels", 20, False, RGBColor(0xCF,0xDA,0xE8), HEAD, True)],
     [("Le cas de l'intégration de nouveaux outils numériques au sein d'Arla Groupe", 15.5, False, RGBColor(0x9FB0C4 >> 16 & 255, 0x9FB0C4 >> 8 & 255, 0x9FB0C4 & 255), BODY, True)]],
    space_after=4)
txt(s, 0.9, 6.72, 8.6, 0.6,
    [[("Auteur : ", 12.5, True, WHITE, BODY), ("Maxime LOUSTALET", 12.5, False, RGBColor(0xCF,0xDA,0xE8), BODY),
      ("      Directeur de recherche : ", 12.5, True, WHITE, BODY), ("Victor COMBES", 12.5, False, RGBColor(0xCF,0xDA,0xE8), BODY)]],
    space_after=0)
notes(s, (
"[0:00 – ACCROCHE / 45s] Bonjour à toutes et à tous. Madame, Monsieur les membres du jury, "
"je vous remercie de votre présence. Je m'appelle Maxime Loustalet et je vais vous présenter "
"mon mémoire de fin d'études, réalisé sous la direction de Victor Combes.\n\n"
"Ce travail porte sur l'intégration de l'intelligence artificielle dans les PME : entre opportunités "
"perçues et freins organisationnels. Je l'ai construit à partir d'un cas concret, celui d'Arla Groupe, "
"la PME du bâtiment dans laquelle j'effectue mon alternance.\n\n"
"CONSEIL : restez debout, souriez, regardez le jury, ne lisez pas la diapo. Annoncez la durée : "
"« Ma présentation durera une vingtaine de minutes, je resterai ensuite à votre disposition pour vos questions. »"
))

# ================================================================ 2. PLAN
s = slide()
header(s, "Fil conducteur de la présentation", "Plan de la soutenance")
plan = [
    ("01", "Contexte & problématique", "Un enjeu d'actualité, un terrain, une question de recherche"),
    ("02", "Cadre théorique & propositions", "PME, culture, IA : ce que dit la littérature"),
    ("03", "Méthodologie empirique", "Étude qualitative — 8 entretiens semi-directifs"),
    ("04", "Résultats & analyse", "Trois questions de recherche confrontées au terrain"),
    ("05", "Apports & préconisations", "Contributions scientifiques et recommandations managériales"),
    ("06", "Conclusion & limites", "Réponse, apport professionnel et perspectives"),
]
yx = 1.75
for i, (n, t, d) in enumerate(plan):
    col = 0 if i < 3 else 1
    row = i % 3
    x = 0.75 + col*6.15
    y = yx + row*1.55
    card(s, x, y, 5.75, 1.32)
    rect(s, x, y, 0.12, 1.32, ACCENT if col==0 else STEEL)
    txt(s, x+0.28, y+0.20, 1.1, 0.9, [[(n, 34, True, RGBColor(0xE1,0xE7,0xEF), HEAD)]], space_after=0)
    txt(s, x+1.25, y+0.20, 4.3, 0.5, [[(t, 16.5, True, NAVY, HEAD)]], space_after=0)
    txt(s, x+1.25, y+0.72, 4.35, 0.5, [[(d, 11.5, False, TEXTG, BODY)]], space_after=0)
footer(s, 2)
notes(s, (
"[0:45 – PLAN / 45s] Ma présentation suivra six temps.\n\n"
"1) Je poserai le CONTEXTE et la problématique. 2) J'exposerai le CADRE THÉORIQUE et les propositions "
"qui en découlent. 3) Je présenterai ma MÉTHODOLOGIE empirique. 4) J'analyserai les RÉSULTATS "
"autour de mes trois questions de recherche. 5) J'en tirerai les APPORTS scientifiques et mes "
"PRÉCONISATIONS managériales. 6) Je conclurai sur l'apport professionnel, les limites et les perspectives.\n\n"
"ASTUCE : ce plan sert de « carte » — le jury doit savoir où vous l'emmenez. Ne vous y attardez pas."
))

# ================================================================ 3. DIVIDER 01
def divider(num, title, subtitle, page):
    s = slide()
    bg(s, NAVY)
    rect(s, 0, 0, 0.28, 7.5, ACCENT)
    rect(s, 0.75, 2.35, 2.4, 0.06, ACCENT)
    txt(s, 0.75, 2.55, 11, 1.4, [[(num, 92, True, RGBColor(0x24,0x46,0x6B), HEAD)]], space_after=0)
    txt(s, 3.0, 2.75, 9.4, 1.6, [[(title, 40, True, WHITE, HEAD)]], space_after=0, anchor=MSO_ANCHOR.TOP)
    txt(s, 3.05, 3.95, 9.2, 0.8, [[(subtitle, 16, False, RGBColor(0xB8,0xC6,0xD8), BODY, True)]], space_after=0)
    footer(s, page, dark=True)
    return s

s = divider("01", "Contexte & problématique", "Un enjeu d'actualité majeur, un terrain concret, une question de recherche", 3)
notes(s, (
"[1:30 – TRANSITION / 10s] Commençons par le contexte et la problématique.\n\n"
"Les diapositives de section servent de respiration : annoncez simplement « Premier temps : le contexte et la problématique » et enchaînez."
))

# ================================================================ 4. CONTEXTE / ACCROCHE
s = slide()
header(s, "01 · Contexte & actualité", "Une bascule technologique, un angle mort : la PME")
# left column - IA montante
card(s, 0.75, 1.75, 5.75, 4.9, CARD_ALT, line=None)
txt(s, 1.05, 1.98, 5.2, 0.5, [[("UNE VAGUE TECHNOLOGIQUE INÉDITE", 13, True, STEEL, HEAD)]], space_after=0)
bullets(s, 1.05, 2.55, 5.2, 4.0, [
    ("ChatGPT, Copilot, Gamma…", " l'IA générative s'installe dans le travail quotidien et modifie les usages."),
    ("Industrie 4.0", " : automatisation, données, IA — la continuité des révolutions industrielles (Rouas & Al Meriouh, 2025)."),
    ("D'ici 2030", " : l'IA pourrait représenter ≈ 14 % du PIB mondial, soit ≈ 15 000 Md$ (Rouas & Al Meriouh, 2025)."),
    ("Un levier stratégique", " de productivité, d'aide à la décision et de compétitivité."),
], size=14, gap=11)
# right column - la PME
card(s, 6.8, 1.75, 5.75, 4.9, CARD, line=LINE_G)
rect(s, 6.8, 1.75, 5.75, 0.12, ACCENT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
txt(s, 7.1, 1.98, 5.2, 0.5, [[("MAIS LA PME EST UN CAS À PART", 13, True, ACCENT_D, HEAD)]], space_after=0)
bullets(s, 7.1, 2.55, 5.2, 4.0, [
    ("≈ 90 % des entreprises mondiales", " et moteur de l'emploi et des territoires (Farsad, 2021)."),
    ("Ressources limitées", " : financières, humaines et technologiques."),
    ("Dirigeant central", " et structure souple (Nassou & Bennani, 2024)."),
    ("Poids des traditions", " et culture d'entreprise (Schein, 2015 ; Torrès, 1999)."),
], size=14, gap=11, marker_color=ACCENT)
footer(s, 4)
notes(s, (
"[1:40 – CONTEXTE / 1min30] Pourquoi ce sujet, et pourquoi maintenant ?\n\n"
"D'un côté, nous vivons une bascule technologique inédite. L'IA générative — ChatGPT, Copilot, Gamma — "
"est entrée dans le quotidien de travail en moins de deux ans. Elle s'inscrit dans le mouvement de "
"l'Industrie 4.0, dans la continuité des révolutions industrielles. Les chiffres donnent le vertige : "
"selon Rouas et Al Meriouh (2025), d'ici 2030 l'IA pourrait peser près de 14 % du PIB mondial, environ "
"15 000 milliards de dollars. C'est devenu un levier stratégique de productivité et de compétitivité.\n\n"
"MAIS — et c'est là mon angle — la PME est un cas particulier. Elle représente près de 90 % des entreprises "
"dans le monde et reste le moteur de l'emploi et des territoires. Or elle dispose de ressources limitées, "
"son dirigeant est au cœur des décisions, sa structure est souple, et elle est fortement marquée par ses "
"traditions et sa culture. L'IA y arrive donc dans des conditions très différentes de celles d'un grand groupe.\n\n"
"C'est cette TENSION — opportunités perçues d'un côté, freins organisationnels de l'autre — qui fonde mon travail."
))

# ================================================================ 5. TERRAIN + PROBLÉMATIQUE
s = slide()
header(s, "01 · Terrain & problématique", "Arla Groupe : un usage réel de l'IA, mais non encadré")
# terrain card
card(s, 0.75, 1.72, 5.15, 4.95, CARD, line=LINE_G)
rect(s, 0.75, 1.72, 5.15, 0.12, STEEL, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
txt(s, 1.02, 1.95, 4.7, 0.5, [[("LE TERRAIN — ARLA GROUPE", 13, True, STEEL, HEAD)]], space_after=0)
bullets(s, 1.02, 2.55, 4.65, 4.0, [
    ("PME du bâtiment", " : maçonnerie, construction métallique, bureau d'études."),
    ("Métiers variés", " : dirigeants, métreurs-deviseurs, conducteurs de travaux, commerciaux, RH, terrain."),
    ("Déjà des usages IA", " : ChatGPT, Copilot, Gamma dans les missions quotidiennes."),
    ("Le constat", " : un usage individuel, spontané, peu encadré et non partagé."),
], size=13.5, gap=10)
# problématique highlight
card(s, 6.15, 1.72, 6.4, 3.05, NAVY, line=None)
rect(s, 6.15, 1.72, 0.14, 3.05, ACCENT)
txt(s, 6.55, 1.98, 5.8, 0.4, [[("PROBLÉMATIQUE", 13, True, ACCENT, HEAD)]], space_after=0)
txt(s, 6.55, 2.45, 5.75, 2.3,
    [[("« Comment les PME peuvent-elles intégrer l'IA et les technologies numériques tout en tenant compte de leurs ressources limitées, de leur structure organisationnelle et de leur culture d'entreprise ? »",
       16.5, True, WHITE, HEAD, True)]], space_after=0, line_spacing=1.05)
# 3 QR chips
qr = [("QR1", "Comment les salariés perçoivent-ils l'IA ?"),
      ("QR2", "Quels freins (organisationnels, humains, culturels) ?"),
      ("QR3", "En quoi les spécificités PME influencent l'adoption ?")]
for i,(a,b) in enumerate(qr):
    y = 4.95 + i*0.57
    rect(s, 6.15, y, 0.9, 0.46, STEEL, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    txt(s, 6.15, y+0.075, 0.9, 0.4, [[(a, 12.5, True, WHITE, HEAD)]], align=PP_ALIGN.CENTER, space_after=0)
    txt(s, 7.2, y+0.055, 5.3, 0.45, [[(b, 12.5, False, TEXTD, BODY)]], space_after=0, anchor=MSO_ANCHOR.MIDDLE)
footer(s, 5)
notes(s, (
"[3:10 – TERRAIN + PROBLÉMATIQUE / 1min30] Mon terrain, c'est Arla Groupe, une PME du bâtiment : "
"maçonnerie, construction métallique, bureau d'études. On y trouve des métiers très variés — des "
"dirigeants, des métreurs-deviseurs, des conducteurs de travaux, des commerciaux, des RH, et des équipes "
"de terrain.\n\n"
"Le point de départ est un constat concret : plusieurs salariés utilisent DÉJÀ l'IA — ChatGPT, Copilot, "
"Gamma — dans leurs missions. Mais cet usage est individuel, spontané, non encadré et non partagé. "
"Chacun fait « dans son coin ».\n\n"
"De ce constat découle ma problématique [la lire lentement, en la montrant] : « Comment les PME "
"peuvent-elles intégrer l'IA et les technologies numériques tout en tenant compte de leurs ressources "
"limitées, de leur structure organisationnelle et de leur culture d'entreprise ? »\n\n"
"Pour y répondre, je l'ai déclinée en TROIS questions de recherche : la perception de l'IA, les freins "
"rencontrés, et le rôle des spécificités propres aux PME."
))

# ================================================================ 6. DIVIDER 02
s = divider("02", "Cadre théorique & propositions", "Ce que la littérature nous apprend sur les PME et sur l'IA", 6)
notes(s, (
"[4:40 – TRANSITION / 10s] J'en viens à mon cadre théorique, en deux volets : d'abord la PME "
"comme objet organisationnel, ensuite l'IA comme levier et défi. J'en tirerai des propositions de recherche."
))

# ================================================================ 7. CADRE 1 - PME
s = slide()
header(s, "02 · Cadre théorique (1/2)", "La PME : un objet organisationnel spécifique")
cols = [
    ("DÉFINITION", STEEL, [
        ("Critères UE", " : < 250 salariés, CA < 50 M€ (Oriot & Misiaszek, 2012)."),
        ("Structure simple", " et peu formalisée (Nassou & Bennani, 2024)."),
    ]),
    ("PROXIMITÉ", ACCENT, [
        ("Dirigeant impliqué", " au quotidien, décisions rapides et centralisées."),
        ("« Proximité organisationnelle »", " : atout relationnel clé (Torrès, 1999)."),
    ]),
    ("CULTURE & TRADITIONS", NAVY2, [
        ("Culture = valeurs partagées", " qui guident les décisions (Schein, 2015)."),
        ("Traditions de métier", " : artisan de métier vs entrepreneurial (Blanchard & Albert-Cromarias, 2022)."),
    ]),
]
for i,(t,c,items) in enumerate(cols):
    x = 0.75 + i*4.08
    card(s, x, 1.75, 3.78, 3.55, CARD, line=LINE_G)
    rect(s, x, 1.75, 3.78, 0.5, c, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    rect(s, x, 2.1, 3.78, 0.18, c)
    txt(s, x+0.2, 1.83, 3.5, 0.4, [[(t, 12.5, True, WHITE, HEAD)]], space_after=0)
    bullets(s, x+0.22, 2.5, 3.42, 2.7, items, size=12.5, gap=9, marker_color=c)
# proposition band
rect(s, 0.75, 5.6, 11.83, 1.0, LIGHT)
rect(s, 0.75, 5.6, 0.14, 1.0, ACCENT)
txt(s, 1.05, 5.72, 1.9, 0.4, [[("PROPOSITION 1", 12.5, True, ACCENT_D, HEAD)]], space_after=0)
txt(s, 3.0, 5.7, 9.4, 0.85,
    [[("Les caractéristiques structurelles et culturelles de la PME (dirigeant, proximité, traditions) façonnent la manière dont elle accueille — ou freine — une innovation comme l'IA.",
       13.5, False, TEXTD, BODY, True)]], space_after=0, anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.03)
footer(s, 7)
notes(s, (
"[4:50 – CADRE PME / 2min] Premier volet : qu'est-ce qu'une PME, du point de vue organisationnel ?\n\n"
"Sur le plan de la DÉFINITION, on retient les critères européens rappelés par Oriot et Misiaszek (2012) : "
"moins de 250 salariés, un chiffre d'affaires sous 50 millions d'euros. Mais au-delà des chiffres, "
"Nassou et Bennani (2024) montrent que la PME a une structure simple, peu formalisée.\n\n"
"Deuxième idée forte : la PROXIMITÉ. Le dirigeant est impliqué dans le quotidien, les décisions sont "
"rapides et centralisées. Torrès (1999) parle de « proximité organisationnelle » : c'est un atout "
"relationnel majeur, une caractéristique qui définit la PME.\n\n"
"Troisième idée : la CULTURE et les TRADITIONS. Schein (2015) définit la culture comme un ensemble de "
"valeurs partagées qui orientent les décisions et la manière de faire face au changement. Et dans les "
"métiers artisanaux, Blanchard et Albert-Cromarias (2022) distinguent l'artisan « de métier », attaché "
"à son savoir-faire, et l'artisan « entrepreneurial », tourné vers la gestion et la croissance.\n\n"
"J'en tire une première PROPOSITION : ces spécificités — dirigeant, proximité, traditions — façonnent "
"la façon dont la PME accueille ou freine une innovation comme l'IA."
))

# ================================================================ 8. CADRE 2 - IA
s = slide()
header(s, "02 · Cadre théorique (2/2)", "L'IA : un levier de performance et un défi organisationnel")
# Opportunités
card(s, 0.75, 1.75, 5.75, 3.55, CARD, line=LINE_G)
rect(s, 0.75, 1.75, 5.75, 0.12, RGBColor(0x2E,0x8B,0x57), shape=MSO_SHAPE.ROUNDED_RECTANGLE)
txt(s, 1.05, 1.98, 5.2, 0.4, [[("▲  OPPORTUNITÉS", 14, True, RGBColor(0x25,0x74,0x48), HEAD)]], space_after=0)
bullets(s, 1.05, 2.55, 5.2, 2.7, [
    ("Transformation digitale → performance", " organisationnelle (Radoui & Cherradi, 2025)."),
    ("Analyse prédictive & agilité", " : anticiper la demande, ajuster la stratégie (Bennour & Oukassi, 2025)."),
    ("Automatisation", " des tâches répétitives à faible valeur ajoutée."),
], size=13.5, gap=10, marker_color=RGBColor(0x2E,0x8B,0x57))
# Freins/menaces
card(s, 6.8, 1.75, 5.75, 3.55, CARD, line=LINE_G)
rect(s, 6.8, 1.75, 5.75, 0.12, RGBColor(0xC0,0x53,0x3B), shape=MSO_SHAPE.ROUNDED_RECTANGLE)
txt(s, 7.1, 1.98, 5.2, 0.4, [[("▼  FREINS & MENACES", 14, True, RGBColor(0xA8,0x45,0x30), HEAD)]], space_after=0)
bullets(s, 7.1, 2.55, 5.2, 2.7, [
    ("Qualité des données, compétences, intégration", " au SI existant (Kokina et al., 2025 ; Duarte, 2025)."),
    ("Enjeux éthiques", " : biais, opacité des algorithmes, confidentialité."),
    ("Dépendance cognitive", " : l'IA amplifie les vulnérabilités humaines (Soro Torna, 2024)."),
], size=13.5, gap=10, marker_color=RGBColor(0xC0,0x53,0x3B))
# proposition band
rect(s, 0.75, 5.6, 11.83, 1.0, LIGHT)
rect(s, 0.75, 5.6, 0.14, 1.0, ACCENT)
txt(s, 1.05, 5.72, 1.9, 0.4, [[("PROPOSITION 2", 12.5, True, ACCENT_D, HEAD)]], space_after=0)
txt(s, 3.0, 5.7, 9.4, 0.85,
    [[("L'adoption de l'IA ne dépend pas que de la technologie : elle se joue autant sur des facteurs humains, organisationnels et culturels que sur des facteurs techniques.",
       13.5, False, TEXTD, BODY, True)]], space_after=0, anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.03)
footer(s, 8)
notes(s, (
"[6:50 – CADRE IA / 2min] Deuxième volet : l'IA elle-même, telle que la traite la littérature. "
"Elle a deux visages.\n\n"
"Côté OPPORTUNITÉS : Radoui et Cherradi (2025) montrent que la transformation digitale améliore la "
"performance organisationnelle sur les plans financier, économique et social. Bennour et Oukassi (2025) "
"insistent sur l'analyse prédictive et l'agilité : anticiper la demande, ajuster la stratégie. Et bien sûr "
"l'automatisation des tâches répétitives à faible valeur ajoutée.\n\n"
"Côté FREINS et MENACES : Kokina et al. (2025) et Duarte (2025) pointent la qualité des données, le "
"manque de compétences et la difficulté d'intégration au système d'information existant. S'y ajoutent des "
"enjeux éthiques — biais, opacité des algorithmes, confidentialité. Et une analyse que je trouve "
"particulièrement fine, celle de Soro Torna (2024) : le vrai risque de l'IA serait humain — dépendance "
"cognitive, perte d'autonomie. L'IA agirait comme un « amplificateur des vulnérabilités humaines ».\n\n"
"D'où ma deuxième PROPOSITION : l'adoption de l'IA ne se joue pas que sur la technologie, mais autant "
"sur des facteurs humains, organisationnels et culturels. C'est exactement ce que mon étude va tester sur le terrain."
))

# ================================================================ 9. SYNTHÈSE + PROPOSITIONS
s = slide()
header(s, "02 · Synthèse", "De la théorie aux propositions de recherche")
# central tension
txt(s, 0.75, 1.68, 11.8, 0.5, [[("La PME face à l'IA : une double dynamique", 16, True, NAVY, HEAD)]], space_after=0)
card(s, 0.75, 2.25, 5.75, 1.5, RGBColor(0xEC,0xF4,0xEC), line=None)
txt(s, 1.05, 2.42, 5.2, 0.4, [[("LEVIERS", 12.5, True, RGBColor(0x25,0x74,0x48), HEAD)]], space_after=0)
txt(s, 1.05, 2.82, 5.25, 0.9, [[("Flexibilité · proximité relationnelle · réactivité · dirigeant moteur",
    13.5, False, TEXTD, BODY)]], space_after=0)
card(s, 6.8, 2.25, 5.75, 1.5, RGBColor(0xF7,0xEC,0xE9), line=None)
txt(s, 7.1, 2.42, 5.2, 0.4, [[("FREINS", 12.5, True, RGBColor(0xA8,0x45,0x30), HEAD)]], space_after=0)
txt(s, 7.1, 2.82, 5.25, 0.9, [[("Ressources limitées · compétences · traditions · absence de cadre commun",
    13.5, False, TEXTD, BODY)]], space_after=0)
# 3 propositions -> QR
txt(s, 0.75, 3.95, 11.8, 0.4, [[("Trois propositions confrontées au terrain", 16, True, NAVY, HEAD)]], space_after=0)
props = [
    ("QR1", "La perception de l'IA dépend surtout du métier et de l'usage réel — pas seulement de l'âge."),
    ("QR2", "Les freins sont pluriels (humains, financiers, organisationnels) et varient selon les profils."),
    ("QR3", "Les spécificités PME (dirigeant, proximité, traditions) conditionnent l'adoption de l'IA."),
]
for i,(a,b) in enumerate(props):
    y = 4.5 + i*0.72
    rect(s, 0.75, y, 1.0, 0.58, NAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    txt(s, 0.75, y+0.13, 1.0, 0.4, [[(a, 14, True, ACCENT, HEAD)]], align=PP_ALIGN.CENTER, space_after=0)
    card(s, 1.95, y, 10.6, 0.58, LIGHT, line=None)
    txt(s, 2.2, y+0.055, 10.2, 0.5, [[(b, 13.5, False, TEXTD, BODY)]], space_after=0, anchor=MSO_ANCHOR.MIDDLE)
footer(s, 9)
notes(s, (
"[8:50 – SYNTHÈSE / 1min] En synthèse, la littérature dessine une DOUBLE dynamique.\n\n"
"D'un côté, la PME a des LEVIERS : sa flexibilité, sa proximité relationnelle, sa réactivité, un dirigeant "
"qui peut jouer un rôle moteur. De l'autre, des FREINS : des ressources limitées, un déficit de compétences, "
"le poids des traditions, et souvent l'absence d'un cadre commun.\n\n"
"De cette tension, je tire trois PROPOSITIONS, une par question de recherche. Pour la QR1, je fais "
"l'hypothèse que la perception de l'IA dépend surtout du métier et de l'usage réel, pas seulement de l'âge. "
"Pour la QR2, que les freins sont pluriels et varient selon les profils. Pour la QR3, que les spécificités "
"de la PME — dirigeant, proximité, traditions — conditionnent l'adoption.\n\n"
"Ce sont ces trois propositions que mon étude de terrain va confronter à la réalité."
))

# ================================================================ 10. DIVIDER 03
s = divider("03", "Méthodologie empirique", "Une étude qualitative exploratoire fondée sur 8 entretiens semi-directifs", 10)
notes(s, (
"[9:50 – TRANSITION / 10s] Voyons maintenant comment j'ai construit mon étude de terrain."
))

# ================================================================ 11. MÉTHODOLOGIE
s = slide()
header(s, "03 · Démarche empirique", "Une étude qualitative exploratoire assumée")
# left: choix + collecte
card(s, 0.75, 1.75, 5.75, 4.9, CARD, line=LINE_G)
rect(s, 0.75, 1.75, 5.75, 0.12, STEEL, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
txt(s, 1.05, 1.98, 5.2, 0.4, [[("POURQUOI LE QUALITATIF ?", 13, True, STEEL, HEAD)]], space_after=0)
bullets(s, 1.05, 2.5, 5.2, 4.0, [
    ("Explorer le « pourquoi » et le « comment »", " des comportements, pas les mesurer (Jando, 2024)."),
    ("8 entretiens semi-directifs", " · 10 questions ouvertes · enregistrés."),
    ("Analyse thématique", " : ≈ 5 h de transcription (Annexe 2) + extraction de verbatims (Annexe 3)."),
    ("Confronter la théorie au terrain", " et faire émerger des facteurs non anticipés."),
], size=13.5, gap=11)
# right: échantillon
card(s, 6.8, 1.75, 5.75, 4.9, CARD, line=LINE_G)
rect(s, 6.8, 1.75, 5.75, 0.12, ACCENT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
txt(s, 7.1, 1.98, 5.2, 0.4, [[("UN ÉCHANTILLONNAGE RAISONNÉ", 13, True, ACCENT_D, HEAD)]], space_after=0)
txt(s, 7.1, 2.5, 5.2, 0.35, [[("3 familles de profils", 13.5, True, NAVY, HEAD)]], space_after=0)
for i,(a) in enumerate(["Dirigeants / responsables — vision stratégique",
                        "Bureau : commercial, administratif, RH — usage concret",
                        "Terrain : chefs d'équipe, ouvriers — métiers moins digitalisés"]):
    y=2.9+i*0.42
    rect(s, 7.1, y+0.07, 0.14, 0.14, ACCENT, shape=MSO_SHAPE.OVAL)
    txt(s, 7.35, y, 5.0, 0.4, [[(a, 12, False, TEXTG, BODY)]], space_after=0)
txt(s, 7.1, 4.35, 5.2, 0.35, [[("3 tranches d'âge", 13.5, True, NAVY, HEAD)]], space_after=0)
for i,(a) in enumerate(["18–29 ans : à l'aise avec le numérique",
                        "30–44 ans : profils intermédiaires, usages mixtes",
                        "45 ans et + : habitudes de travail installées"]):
    y=4.75+i*0.42
    rect(s, 7.1, y+0.07, 0.14, 0.14, STEEL, shape=MSO_SHAPE.OVAL)
    txt(s, 7.35, y, 5.0, 0.4, [[(a, 12, False, TEXTG, BODY)]], space_after=0)
# forces/limites strip
rect(s, 0.75, 6.75, 11.83, 0.02, LINE_G)
footer(s, 11)
notes(s, (
"[10:00 – MÉTHODOLOGIE / 2min] J'ai choisi une approche QUALITATIVE EXPLORATOIRE. Pourquoi ? Parce que "
"ma problématique porte sur des perceptions, des freins, des dynamiques humaines. Comme le dit Jando (2024), "
"le qualitatif permet d'explorer le « pourquoi » et le « comment » des comportements — ce qu'aucun chiffre "
"ne donne.\n\n"
"Concrètement : j'ai mené 8 entretiens semi-directifs, avec une trame de 10 questions ouvertes, tous "
"enregistrés. J'ai ensuite transcrit — environ 5 heures de travail — puis extrait les verbatims les plus "
"significatifs pour les confronter à mes trois questions de recherche.\n\n"
"Le point clé, c'est l'ÉCHANTILLONNAGE RAISONNÉ. J'ai croisé deux dimensions. D'abord trois familles de "
"profils : les dirigeants pour la vision stratégique ; le bureau — commercial, administratif, RH — pour "
"l'usage concret ; et le terrain — chefs d'équipe, ouvriers — pour les métiers les moins digitalisés. "
"Ensuite trois tranches d'âge, pour capter d'éventuelles différences générationnelles.\n\n"
"FORCES : on recueille un avis « à chaud », en profondeur. LIMITES, que j'assume : échantillon restreint, "
"une seule entreprise, résultats non généralisables et une part de subjectivité. J'y reviendrai en conclusion."
))

# ================================================================ 12. DIVIDER 04
s = divider("04", "Résultats & analyse", "Trois questions de recherche confrontées aux 8 entretiens", 12)
notes(s, (
"[12:00 – TRANSITION / 10s] J'en arrive au cœur de mon travail : les résultats, question de recherche par "
"question de recherche, en dialogue constant avec la littérature."
))

# ================================================================ 13. QR1
def result_slide(kicker, title, page, headline, items, litt):
    s = slide()
    header(s, kicker, title)
    card(s, 0.75, 1.72, 11.83, 0.72, NAVY, line=None)
    rect(s, 0.75, 1.72, 0.14, 0.72, ACCENT)
    txt(s, 1.05, 1.79, 11.2, 0.6, [[("Résultat clé — ", 14.5, True, ACCENT, HEAD),
        (headline, 14.5, False, WHITE, BODY, True)]], space_after=0, anchor=MSO_ANCHOR.MIDDLE)
    bullets(s, 0.9, 2.75, 11.6, 3.2, items, size=13.8, gap=10)
    rect(s, 0.75, 6.02, 11.83, 0.62, CARD_ALT, line=None)
    rect(s, 0.75, 6.02, 0.14, 0.62, STEEL)
    txt(s, 1.05, 6.09, 11.3, 0.5, [[("Lien littérature — ", 12.5, True, STEEL, HEAD),
        (litt, 12.5, False, TEXTD, BODY, True)]], space_after=0, anchor=MSO_ANCHOR.MIDDLE)
    footer(s, page)
    return s

s = result_slide("04 · Question de recherche 1", "Comment les salariés perçoivent-ils l'IA ?", 13,
    "une perception globalement positive, mais un usage isolé et non concerté.",
    [
        ("Le gain de temps, argument unanime :", " l'IA « recentre sur les tâches à valeur ajoutée » ; la directrice commerciale veut « arrêter la bureautique basique »."),
        ("Les non-utilisateurs ne rejettent pas :", " l'ouvrier (57 ans) et le dirigeant (63 ans) ne connaissent pas l'IA, mais parlent d'un « manque d'utilité », pas d'un refus."),
        ("Une lecture par le métier :", " bureau automatisable vs terrain — « le béton ne se verse pas de lui-même »."),
        ("L'âge ne suffit pas :", " « l'âge n'est pas le seul facteur… certains jeunes restent attachés à leurs habitudes »."),
        ("Un usage individuel :", " ChatGPT, Copilot, Gamma utilisés sans savoir « jusqu'où on a le droit » — « c'est une affaire personnelle »."),
    ],
    "valide Bennour & Oukassi (2025) sur la productivité ; illustre le défaut de structure collective pointé par Kokina et al. (2025).")
notes(s, (
"[12:10 – QR1 PERCEPTION / 2min] Première question : comment les salariés perçoivent-ils l'IA ?\n\n"
"Le résultat clé : la perception est globalement POSITIVE, mais l'usage reste isolé.\n\n"
"Premier point, unanime : le GAIN DE TEMPS. Tous ceux qui l'utilisent y voient un moyen de se recentrer "
"sur les tâches à valeur ajoutée. La directrice commerciale parle d'« arrêter de faire de la bureautique "
"basique ». Cela valide directement Bennour et Oukassi (2025).\n\n"
"Deuxième point, plus surprenant : les non-utilisateurs — un ouvrier de 57 ans, un dirigeant de 63 ans — "
"ne REJETTENT pas l'IA. Ils disent ne pas la connaître, ne pas en voir l'utilité dans leur métier. "
"Ce n'est pas de l'hostilité.\n\n"
"Troisième point : la vraie ligne de partage, c'est le MÉTIER. Le travail de bureau est automatisable, "
"le terrain beaucoup moins — je cite l'ouvrier : « le béton ne se verse pas de lui-même ».\n\n"
"Quatrième point : l'âge joue, mais ne suffit pas. Un cadre observe que « certains jeunes restent très "
"attachés à leurs habitudes ». Cela nuance la lecture purement générationnelle.\n\n"
"Enfin, l'usage est INDIVIDUEL et non concerté : chacun utilise son outil sans savoir jusqu'où il a le "
"droit — « c'est une affaire personnelle ». C'est exactement le déficit de structure collective décrit "
"par Kokina et al. (2025). Ma première proposition est donc confirmée, avec une nuance : le métier compte "
"encore plus que l'âge."
))

# ================================================================ 14. QR2
s = result_slide("04 · Question de recherche 2", "Quels sont les freins à l'adoption de l'IA ?", 14,
    "des freins pluriels ; le frein prioritaire n'est pas le même selon les profils.",
    [
        ("Deux freins partagés par tous :", " le besoin de formation / accompagnement (unanime) et la contrainte financière (licences, développements sur mesure)."),
        ("Un frein propre à chaque profil :", " appréhension psychologique (RH), manque d'esprit critique face aux chiffres (achat), confidentialité des données chantier (conducteur)."),
        ("Un frein émergent, absent de la littérature :", " la peur d'une perte de compétences par sur-délégation à l'IA (métreur-deviseur)."),
        ("La résistance = ancrage des habitudes :", " « 40 ans qu'on travaille d'une certaine façon » — mais c'est une vigilance professionnelle, pas un rejet idéologique."),
        ("Un frein structurel, le plus facile à lever :", " l'absence de politique commune et de référent interne."),
    ],
    "confirme Kokina et al. (2025) & Duarte (2025) sur les compétences ; l'ancrage renvoie à Schein (2015) et Blanchard & Albert-Cromarias (2022).")
notes(s, (
"[14:10 – QR2 FREINS / 2min] Deuxième question : quels sont les freins ?\n\n"
"Le résultat le plus marquant, c'est que le frein PRIORITAIRE n'est pas le même selon les profils — là où "
"la littérature a tendance à homogénéiser les obstacles.\n\n"
"Deux freins sont toutefois PARTAGÉS par tous : le besoin de formation et d'accompagnement — unanime, y "
"compris chez ceux qui n'utilisent pas l'IA — et la contrainte financière, qu'il s'agisse des licences ou "
"des développements sur mesure. Cela confirme Kokina et al. (2025) et Duarte (2025) sur le déficit de "
"compétences.\n\n"
"Ensuite, chaque profil a SON frein prioritaire : pour la RH, c'est l'appréhension psychologique, « la peur "
"de ne pas être capable de maîtriser » ; pour l'assistant achat, le manque d'esprit critique face aux "
"réponses chiffrées ; pour le conducteur de travaux, la confidentialité des données chantier.\n\n"
"Un frein m'a particulièrement marqué car il est ABSENT de la littérature que j'ai mobilisée : le "
"métreur-deviseur redoute une PERTE DE COMPÉTENCES à force de trop déléguer à l'IA. J'y reviendrai "
"dans mes apports.\n\n"
"La résistance au changement existe — « 40 ans qu'on travaille d'une certaine façon » — et renvoie à "
"l'ancrage décrit par Schein (2015). Mais attention : ce n'est pas un rejet idéologique, c'est une "
"vigilance professionnelle.\n\n"
"Enfin, le frein le plus STRUCTUREL — l'absence de politique commune — est aussi le plus facile à lever. "
"C'est le point de départ de mes préconisations."
))

# ================================================================ 15. QR3
s = result_slide("04 · Question de recherche 3", "En quoi les spécificités PME influencent-elles l'adoption ?", 15,
    "le dirigeant est décisif (8/8) ; la proximité et les traditions sont un cadre, pas un mur.",
    [
        ("Le rôle du dirigeant — convergence totale (8/8) :", " « si les dirigeants ne sont pas convaincus, le changement sera plus lent » — un seul leader moteur peut suffire."),
        ("La proximité, levier ET risque :", " elle diffuse les bonnes pratiques (compagnonnage) mais peut aussi propager les résistances (Torrès, 1999)."),
        ("La contrainte financière, différenciée :", " gérable pour le grand public, bloquante sur le sur-mesure ; + un enjeu concurrentiel inédit (E5)."),
        ("Top-down vs participatif :", " la direction impulse, mais le terrain aspire à co-décider — « c'est un référendum, tous ensemble »."),
        ("Les traditions = adoption sélective :", " personne ne rejette l'IA au nom du métier — elle est le contexte de l'adoption, pas un obstacle."),
    ],
    "valide Nassou & Bennani (2024) et Torrès (1999) ; complète Blanchard & Albert-Cromarias (2022) par la notion d'« adoption sélective ».")
notes(s, (
"[16:10 – QR3 SPÉCIFICITÉS PME / 2min] Troisième question : en quoi les spécificités de la PME jouent-elles ?\n\n"
"Le résultat le plus net de toute mon étude : le RÔLE DU DIRIGEANT. Sur 8 entretiens, 8 le citent comme "
"facteur déterminant — y compris un ouvrier non impliqué dans la stratégie. Le dirigeant lui-même le "
"reconnaît : « si les dirigeants ne sont pas convaincus, le changement sera beaucoup plus lent. » Et un "
"enseignement fort : un seul leader moteur peut suffire à enclencher la dynamique. Cela valide Nassou et "
"Bennani (2024) et Torrès (1999).\n\n"
"La PROXIMITÉ organisationnelle de Torrès joue double jeu : elle diffuse les bonnes pratiques par "
"compagnonnage, mais elle peut aussi propager les résistances si un profil influent freine.\n\n"
"La contrainte FINANCIÈRE est réelle mais différenciée : gérable pour les outils grand public — ChatGPT "
"gratuit, Copilot dans le Pack Office — bloquante pour le sur-mesure. Un commercial du CODIR ajoute une "
"dimension que je n'ai trouvée nulle part dans la littérature : le risque CONCURRENTIEL pour les PME qui "
"n'investiraient pas.\n\n"
"J'ai aussi observé une tension entre logique TOP-DOWN et aspiration PARTICIPATIVE : un salarié parle d'un "
"« référendum, tous ensemble ». Dans une PME à équipes jeunes, l'autorité seule ne suffit pas.\n\n"
"Enfin, sur les TRADITIONS : personne ne rejette l'IA au nom du métier. Elles sont le contexte de "
"l'adoption, pas un mur. Je propose de parler d'« adoption sélective », ce qui complète Blanchard et "
"Albert-Cromarias (2022)."
))

# ================================================================ 16. APPORTS SCIENTIFIQUES
s = slide()
header(s, "04 · Apports scientifiques", "Cinq résultats non anticipés par le modèle théorique")
apports = [
    ("Vigilance ≠ refus", "La distance des non-utilisateurs est une vigilance professionnelle temporaire, pas un rejet de l'IA."),
    ("Un seul dirigeant moteur suffit", "Il n'est pas nécessaire de convaincre tout l'encadrement pour enclencher la dynamique."),
    ("Le risque de perte de compétences", "La sur-délégation à l'IA fait craindre une érosion du savoir-faire — frein peu documenté."),
    ("La dimension concurrentielle", "Ne pas investir dans l'IA devient un désavantage compétitif pour la PME."),
    ("L'« adoption sélective »", "La tradition n'est pas un frein : elle est le cadre dans lequel l'adoption se construit, métier par métier."),
]
for i,(t,d) in enumerate(apports):
    col = i % 2
    row = i // 2
    x = 0.75 + col*6.05
    y = 1.78 + row*1.28
    if i == 4:
        x = 0.75; y = 1.78 + 2*1.28; w = 11.83
    else:
        w = 5.75
    card(s, x, y, w, 1.12, CARD, line=LINE_G)
    rect(s, x, y, 0.5, 1.12, NAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    rect(s, x+0.25, y, 0.25, 1.12, NAVY)
    txt(s, x-0.02, y+0.32, 0.5, 0.5, [[(str(i+1), 22, True, ACCENT, HEAD)]], align=PP_ALIGN.CENTER, space_after=0)
    txt(s, x+0.72, y+0.16, w-0.9, 0.4, [[(t, 14.5, True, NAVY, HEAD)]], space_after=0)
    txt(s, x+0.72, y+0.55, w-0.95, 0.55, [[(d, 12, False, TEXTG, BODY)]], space_after=0)
footer(s, 16)
notes(s, (
"[18:10 – APPORTS SCIENTIFIQUES / 1min30] Au-delà de la confirmation de mes propositions, mon terrain a "
"fait émerger CINQ résultats que le modèle théorique n'anticipait pas — c'est la contribution originale "
"de mon travail.\n\n"
"Un : la distance des non-utilisateurs est une VIGILANCE professionnelle temporaire, pas un refus. "
"Deux : un SEUL dirigeant moteur suffit à lancer la dynamique — inutile de convaincre tout le monde d'emblée. "
"Trois : un frein peu documenté, la peur d'une PERTE DE COMPÉTENCES par sur-délégation à l'IA. "
"Quatre : une dimension CONCURRENTIELLE — ne pas investir devient un désavantage compétitif. "
"Cinq : la tradition n'est pas un obstacle mais un CADRE — je parle d'« adoption sélective », "
"métier par métier.\n\n"
"Ce sont ces apports qui justifient l'intérêt scientifique du mémoire, et ils débouchent directement "
"sur mes recommandations."
))

# ================================================================ 17. DIVIDER 05
s = divider("05", "Apports & préconisations", "Trois recommandations opérationnelles pour Arla Groupe", 17)
notes(s, (
"[19:40 – TRANSITION / 10s] Ces résultats m'ont permis de formuler trois préconisations concrètes, "
"directement actionnables par la direction d'Arla Groupe."
))

# ================================================================ 18. PRÉCONISATIONS
s = slide()
header(s, "05 · Préconisations managériales", "Trois leviers pour une intégration durable de l'IA")
precos = [
    ("01", "Structurer une politique commune", "pilotée par un référent IA interne",
     ["1–2 outils communs à tous", "Un coordinateur IA (informaticien déjà recruté)", "Des règles claires de confidentialité"],
     "Répond au frein : absence de cadre commun"),
    ("02", "Déployer une formation différenciée", "par les gains rapides et le compagnonnage",
     ["Cibler d'abord les « quick wins » métier", "Former des ambassadeurs internes", "Tutorat plutôt que théorie externe"],
     "Répond au frein : formation & compétences"),
    ("03", "Adopter une démarche participative", "portée par la direction, co-construite",
     ["Vision claire portée par le dirigeant", "Espace de dialogue avec les équipes", "Respect des identités métier"],
     "Répond au frein : tension top-down / adhésion"),
]
for i,(n,t,sub,li,fr) in enumerate(precos):
    x = 0.75 + i*4.08
    card(s, x, 1.75, 3.78, 4.55, CARD, line=LINE_G)
    rect(s, x, 1.75, 3.78, 0.9, NAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    rect(s, x, 2.3, 3.78, 0.35, NAVY)
    txt(s, x+0.22, 1.85, 1.0, 0.7, [[(n, 30, True, ACCENT, HEAD)]], space_after=0)
    txt(s, x+1.15, 1.86, 2.55, 0.75, [[(t, 13.5, True, WHITE, HEAD)]], space_after=0, anchor=MSO_ANCHOR.MIDDLE)
    txt(s, x+0.22, 2.78, 3.4, 0.5, [[(sub, 11.5, False, STEEL, BODY, True)]], space_after=0)
    bullets(s, x+0.22, 3.4, 3.4, 2.1, li, size=12, gap=9, marker_color=ACCENT, bold_lead=False)
    rect(s, x+0.0, 5.72, 3.78, 0.58, CARD_ALT, line=None)
    txt(s, x+0.22, 5.8, 3.4, 0.45, [[(fr, 10.5, True, ACCENT_D, BODY, True)]], space_after=0, anchor=MSO_ANCHOR.MIDDLE)
footer(s, 18)
notes(s, (
"[19:50 – PRÉCONISATIONS / 2min] Mes trois préconisations répondent chacune à un frein identifié sur le terrain.\n\n"
"PRÉCONISATION 1 : structurer une politique commune, pilotée par un référent IA interne. Concrètement : "
"choisir un ou deux outils communs à tous, nommer un coordinateur — l'informaticien récemment recruté, "
"cité par plusieurs salariés — et définir des règles claires de confidentialité. Cela répond au frein "
"structurel : l'absence de cadre commun.\n\n"
"PRÉCONISATION 2 : une formation DIFFÉRENCIÉE, fondée sur deux exigences remontées du terrain : d'abord "
"démontrer des gains rapides et concrets — les « quick wins » propres à chaque métier — ensuite s'appuyer "
"sur le compagnonnage : identifier les collaborateurs à l'aise, en faire des AMBASSADEURS. C'est peu coûteux "
"et adapté à une PME. Cela répond au frein de la formation et des compétences.\n\n"
"PRÉCONISATION 3 : une démarche PARTICIPATIVE. La direction impulse une vision claire — pourquoi l'IA, "
"pour quels objectifs, à quel rythme — mais co-construit avec les équipes, en ouvrant un espace de dialogue "
"qui respecte les identités métier. Cela résout la tension entre top-down et adhésion.\n\n"
"Ces trois leviers forment un plan cohérent, réaliste, à la portée d'une PME comme Arla Groupe."
))

# ================================================================ 19. APPORT PROFESSIONNEL
s = slide()
header(s, "05 · Apport dans ma vie professionnelle", "Ce que ce mémoire m'apporte, et apporte à l'entreprise")
left = [
    ("Une posture de conduite du changement", "Comprendre qu'une techno ne s'impose pas : elle s'accompagne, s'explique et se co-construit."),
    ("Des compétences d'analyse", "Mener des entretiens, écouter, structurer des verbatims, relier théorie et terrain."),
]
right = [
    ("Un livrable directement actionnable", "Un plan en 3 préconisations que je peux porter concrètement chez Arla Groupe, mon entreprise d'alternance."),
    ("Une vision stratégique du numérique", "Un atout pour mon projet professionnel : accompagner la transformation digitale des organisations."),
]
for col,(items) in enumerate([left,right]):
    for i,(t,d) in enumerate(items):
        x = 0.75 + col*6.05
        y = 1.85 + i*2.35
        card(s, x, y, 5.75, 2.05, CARD, line=LINE_G)
        rect(s, x, y, 0.14, 2.05, ACCENT if col==0 else STEEL)
        txt(s, x+0.35, y+0.28, 5.2, 0.6, [[(t, 15, True, NAVY, HEAD)]], space_after=0)
        txt(s, x+0.35, y+0.95, 5.2, 1.0, [[(d, 12.5, False, TEXTG, BODY)]], space_after=0, line_spacing=1.05)
footer(s, 19)
notes(s, (
"[21:50 – APPORT PROFESSIONNEL / 1min] Je voudrais insister sur ce que ce mémoire m'apporte, "
"personnellement et professionnellement — car c'est un critère d'évaluation à part entière.\n\n"
"D'abord, une vraie posture de CONDUITE DU CHANGEMENT : j'ai compris qu'une technologie ne s'impose pas, "
"elle s'accompagne, s'explique, se co-construit. Ensuite, des compétences d'ANALYSE : conduire des "
"entretiens, écouter vraiment, structurer des verbatims, faire dialoguer la théorie et le terrain.\n\n"
"Surtout, ce travail débouche sur un LIVRABLE directement actionnable : mes trois préconisations, je peux "
"les porter concrètement chez Arla Groupe, où je suis en alternance. Enfin, il me donne une vision "
"stratégique du numérique qui nourrit mon projet professionnel : accompagner la transformation digitale "
"des organisations.\n\n"
"Ce mémoire n'est donc pas un exercice théorique : c'est un outil que je vais réellement utiliser."
))

# ================================================================ 20. CONCLUSION
s = slide()
header(s, "06 · Conclusion & perspectives", "Ce qu'il faut retenir")
# réponse
card(s, 0.75, 1.72, 11.83, 1.35, NAVY, line=None)
rect(s, 0.75, 1.72, 0.14, 1.35, ACCENT)
txt(s, 1.1, 1.86, 2.4, 0.4, [[("LA RÉPONSE", 12.5, True, ACCENT, HEAD)]], space_after=0)
txt(s, 1.1, 2.24, 11.2, 0.8, [[("L'adoption de l'IA en PME dépend autant de facteurs humains qu'organisationnels. Les spécificités de la PME — dirigeant, proximité, traditions — ne sont pas un frein insurmontable, mais le ",
    13.5, False, WHITE, BODY), ("cadre à prendre en compte", 13.5, True, WHITE, BODY), (" pour réussir l'intégration.", 13.5, False, WHITE, BODY)]],
    space_after=0, line_spacing=1.05)
# facteurs clés + limites + ouverture
card(s, 0.75, 3.25, 3.78, 3.15, LIGHT, line=None)
txt(s, 1.0, 3.42, 3.4, 0.4, [[("FACTEURS CLÉS", 12.5, True, NAVY, HEAD)]], space_after=0)
bullets(s, 1.0, 3.9, 3.4, 2.4, ["Rôle du dirigeant","Proximité entre salariés","Contraintes financières","Identités professionnelles"], size=12.5, gap=8)
card(s, 4.78, 3.25, 3.78, 3.15, LIGHT, line=None)
txt(s, 5.03, 3.42, 3.4, 0.4, [[("LIMITES", 12.5, True, RGBColor(0xA8,0x45,0x30), HEAD)]], space_after=0)
bullets(s, 5.03, 3.9, 3.4, 2.4, ["8 entretiens, une seule PME","Résultats non généralisables","Part de subjectivité","Instantané dans le temps"], size=12.5, gap=8, marker_color=RGBColor(0xC0,0x53,0x3B))
card(s, 8.8, 3.25, 3.78, 3.15, RGBColor(0xEC,0xF4,0xEC), line=None)
txt(s, 9.05, 3.42, 3.4, 0.4, [[("OUVERTURE", 12.5, True, RGBColor(0x25,0x74,0x48), HEAD)]], space_after=0)
txt(s, 9.05, 3.9, 3.35, 2.4, [[("Prolonger par une étude ",12.5,False,TEXTD,BODY),("quantitative",12.5,True,TEXTD,BODY),(" auprès de tous les salariés, pour hiérarchiser les freins et prioriser les préconisations selon les besoins réels de l'entreprise.",12.5,False,TEXTG,BODY)]], space_after=0, line_spacing=1.08)
footer(s, 20)
notes(s, (
"[22:50 – CONCLUSION / 1min30] Pour conclure, je reviens à ma problématique.\n\n"
"LA RÉPONSE : l'adoption de l'IA dans une PME dépend autant de facteurs HUMAINS qu'ORGANISATIONNELS. Et "
"surtout, les spécificités de la PME — le rôle du dirigeant, la proximité, les traditions — ne sont pas un "
"frein insurmontable : ce sont le CADRE dont il faut tenir compte pour réussir l'intégration.\n\n"
"Quatre FACTEURS CLÉS ressortent : le rôle du dirigeant, la proximité entre salariés, les contraintes "
"financières et la force des identités professionnelles.\n\n"
"Je reste lucide sur les LIMITES : 8 entretiens, une seule entreprise, des résultats non généralisables, "
"une part de subjectivité, et une photographie à un instant donné.\n\n"
"D'où mon OUVERTURE : prolonger ce travail par une étude QUANTITATIVE auprès de l'ensemble des salariés "
"d'Arla Groupe, pour hiérarchiser les freins et prioriser les préconisations selon les besoins réels de "
"l'entreprise.\n\n"
"Je vous remercie de votre attention."
))

# ================================================================ 21. MERCI
s = slide()
bg(s, NAVY)
rect(s, 0, 0, 13.333, 0.28, ACCENT)
rect(s, 0, 7.22, 13.333, 0.28, ACCENT)
txt(s, 0.9, 2.55, 11.5, 1.2, [[("Merci de votre attention", 40, True, WHITE, HEAD)]], space_after=0)
txt(s, 0.9, 3.75, 11.5, 0.7, [[("Je suis à votre disposition pour échanger et répondre à vos questions.", 17, False, RGBColor(0xC6,0xD3,0xE2), BODY, True)]], space_after=0)
rect(s, 0.93, 3.55, 2.2, 0.05, ACCENT)
txt(s, 0.9, 5.4, 11.5, 0.5, [[("Maxime LOUSTALET", 15, True, WHITE, BODY),
    ("   ·   L'intégration de l'IA dans les PME — le cas d'Arla Groupe", 15, False, RGBColor(0x9F,0xB0,0xC4), BODY)]], space_after=0)
txt(s, 0.9, 5.85, 11.5, 0.5, [[("Directeur de recherche : Victor COMBES   ·   Eklore-ed School of Management   ·   2025–2026", 12.5, False, RGBColor(0x7F,0x92,0xAC), BODY)]], space_after=0)
notes(s, (
"[24:20 – CLÔTURE] « Je vous remercie de votre attention et je suis à votre disposition pour vos questions. »\n\n"
"CONSEIL FINAL : marquez un temps d'arrêt, souriez, et laissez le jury ouvrir l'échange. Gardez en tête "
"les questions probables : Pourquoi le qualitatif et pas le quantitatif ? Comment garantir l'objectivité "
"avec 8 entretiens ? Vos préconisations sont-elles chiffrées / calendairisées ? Le ROI de l'IA pour Arla ? "
"Le RGPD et la confidentialité des données chantier ? Comment gérez-vous votre propre position d'alternant "
"(biais d'implication) ? Restez calme, reformulez la question, appuyez-vous sur vos verbatims."
))

prs.save("/home/user/maxime/Soutenance_MFE_LOUSTALET_Maxime.pptx")
print("OK — slides:", len(prs.slides._sldIdLst))
